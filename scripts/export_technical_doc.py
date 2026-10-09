"""Export the English technical and user guides to Word. Requires requirements-docs.txt."""

from pathlib import Path
import re

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE


def inline(paragraph, text):
    for part in re.split(r"(`[^`]+`|\*\*[^*]+\*\*|\[[^\]]+\]\([^)]+\))", text):
        link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part)
        if link:
            relation = paragraph.part.relate_to(link[2], RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
            hyperlink = OxmlElement('w:hyperlink')
            hyperlink.set(qn('r:id'), relation)
            run_element = OxmlElement('w:r')
            properties = OxmlElement('w:rPr')
            style = OxmlElement('w:rStyle')
            style.set(qn('w:val'), 'Hyperlink')
            properties.append(style)
            run_element.append(properties)
            label = OxmlElement('w:t')
            label.text = link[1]
            run_element.append(label)
            hyperlink.append(run_element)
            paragraph._p.append(hyperlink)
            continue
        run = paragraph.add_run(part.strip('`') if part.startswith('`') else
                                part[2:-2] if part.startswith('**') else part)
        if part.startswith('`'):
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
        elif part.startswith('**'):
            run.bold = True


def export(source, destination):
    lines = source.read_text(encoding='utf-8-sig').splitlines()
    title = next((line[2:] for line in lines if line.startswith('# ')), source.stem)
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2)
    section.left_margin = section.right_margin = Cm(1.8)
    normal = document.styles['Normal']
    normal.font.name, normal.font.size = 'Calibri', Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    code_style = document.styles.add_style('Code Block', 1)
    code_style.font.name, code_style.font.size = 'Consolas', Pt(8)
    code_style.paragraph_format.space_after = Pt(0)
    code_style.paragraph_format.line_spacing = 1
    for name in ['Title', 'Heading 1', 'Heading 2', 'Heading 3']:
        document.styles[name].font.color.rgb = RGBColor.from_string('17365D')
    document.core_properties.title = title
    document.core_properties.subject = 'Tytan Updater: retained first Windows distribution and PHP server'
    document.core_properties.language = 'en-US'
    language = OxmlElement('w:lang')
    language.set(qn('w:val'), 'en-US')
    normal.element.get_or_add_rPr().append(language)

    i = 0
    while i < len(lines):
        line = lines[i]
        if line.startswith('```'):
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                paragraph = document.add_paragraph(style='Code Block')
                paragraph.add_run(lines[i] or ' ')
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'F2F4F7')
                paragraph._p.get_or_add_pPr().append(shade)
                i += 1
        elif line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                cells = [cell.strip() for cell in lines[i].strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?', cell) for cell in cells):
                    rows.append(cells)
                i += 1
            table = document.add_table(rows=0, cols=len(rows[0]))
            table.style = 'Light Shading Accent 1'
            for number, row in enumerate(rows):
                cells = table.add_row().cells
                for cell, text in zip(cells, row):
                    inline(cell.paragraphs[0], text)
                    if number == 0:
                        for run in cell.paragraphs[0].runs:
                            run.bold = True
                if number == 0:
                    repeating_header = OxmlElement('w:tblHeader')
                    cells[0]._tc.getparent().get_or_add_trPr().append(repeating_header)
            document.add_paragraph()
            continue
        elif line.startswith('# '):
            document.add_heading(line[2:], 0)
        elif line.startswith('## '):
            document.add_heading(line[3:], 1)
        elif line.startswith('### '):
            document.add_heading(line[4:], 2)
        elif line.startswith('- '):
            inline(document.add_paragraph(style='List Bullet'), line[2:])
        elif re.match(r'^\d+\. ', line):
            inline(document.add_paragraph(style='List Number'), re.sub(r'^\d+\. ', '', line))
        elif line.strip():
            inline(document.add_paragraph(), line)
        i += 1

    header = section.header.paragraphs[0]
    header.text = title.upper()
    header.runs[0].font.size = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.alignment = 2
    footer.add_run('Page ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    document.save(destination)


if __name__ == '__main__':
    root = Path(__file__).resolve().parents[1]
    for name in ('technical-guide', 'user-guide', 'technical-walkthrough'):
        source = root / 'documentation' / (name + '.md')
        destination = source.with_suffix('.docx')
        export(source, destination)
        print(f'Exported {destination.name}')
